from __future__ import annotations

import html
import re
import shutil
import textwrap
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape as xml_escape

from cover_art import CoverMetadata, build_ebook_cover, build_print_cover
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import inch
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch as rl_inch
from reportlab.platypus import (
    BaseDocTemplate,
    Flowable,
    Frame,
    Image as PdfImage,
    ListFlowable,
    ListItem,
    PageBreak,
    PageTemplate,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
)

from print_specs import TRIM_6X9


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "dist"
WORK = ROOT / "work"
MD = ROOT / "manuscrito" / "robo-trade-go-binance-manuscrito.md"
COVER = ROOT / "assets" / "capa-robo-trade-go-binance.jpg"
COVER_THEME = ROOT / "assets" / "cover" / "cover_theme.json"
EPUB = OUT / "robo-trade-go-binance.epub"
DOCX = OUT / "robo-trade-go-binance.docx"
PDF = OUT / "robo-trade-go-binance-revisao.pdf"
PRINT_PDF = OUT / "robo-trade-go-binance-print-6x9.pdf"
PRINT_COVER = OUT / "robo-trade-go-binance-capa-print.pdf"
CODE_DIR = ROOT / "codigo-robo-go-binance"
CODE_ZIP = OUT / "codigo-robo-go-binance.zip"

PRINT_PAPER = "white"

TITLE = "Robôs de Trade com Go e Binance para Iniciantes"
SUBTITLE = (
    "Um guia prático, simples e educativo para criar seu primeiro bot "
    "na Spot Testnet"
)
AUTHOR = "Addo Del Grossi"
COPYRIGHT = "Copyright © 2026 Addo Del Grossi. Todos os direitos reservados."
LANG = "pt-BR"
BACK_BLURB = (
    "Aprenda, passo a passo, a criar um robô educativo de trade com Go e a "
    "Binance Spot Testnet. Você vai instalar o Go, configurar variáveis de "
    "ambiente, buscar preços, ler velas, calcular médias móveis e validar "
    "ordens com segurança — sem operar dinheiro real.\n\n"
    "Um guia direto para iniciantes que querem entender APIs, robôs e o "
    "empacotamento de um livro técnico simples para a Amazon KDP."
)
BACK_DISCLAIMER = (
    "Material educativo. Não é recomendação financeira nem promessa de lucro. "
    "O projeto usa a Binance Spot Testnet."
)


def _cover_metadata() -> CoverMetadata:
    return CoverMetadata(
        title=TITLE,
        subtitle=SUBTITLE,
        author=AUTHOR,
        back_blurb=BACK_BLURB,
        back_disclaimer=BACK_DISCLAIMER,
    )


def make_cover() -> None:
    build_ebook_cover(COVER, _cover_metadata(), COVER_THEME)


def make_print_cover(pages: int) -> None:
    build_print_cover(
        PRINT_COVER,
        COVER,
        pages,
        PRINT_PAPER,
        _cover_metadata(),
        COVER_THEME,
    )


@dataclass
class Block:
    kind: str
    text: str
    level: int = 0
    items: list[str] | None = None
    lang: str = ""


def slugify(text: str) -> str:
    text = text.lower()
    text = re.sub(r"[áàãâä]", "a", text)
    text = re.sub(r"[éèêë]", "e", text)
    text = re.sub(r"[íìîï]", "i", text)
    text = re.sub(r"[óòõôö]", "o", text)
    text = re.sub(r"[úùûü]", "u", text)
    text = text.replace("ç", "c")
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return text or "secao"


def parse_markdown(markdown: str) -> list[Block]:
    blocks: list[Block] = []
    lines = markdown.splitlines()
    i = 0
    paragraph: list[str] = []
    list_kind: str | None = None
    list_items: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph
        if paragraph:
            blocks.append(Block("paragraph", " ".join(paragraph).strip()))
            paragraph = []

    def flush_list() -> None:
        nonlocal list_kind, list_items
        if list_kind and list_items:
            blocks.append(Block(list_kind, "", items=list_items[:]))
        list_kind = None
        list_items = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```") or stripped.startswith("~~~"):
            flush_paragraph()
            flush_list()
            fence = stripped[:3]
            lang = stripped[3:].strip()
            code: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith(fence):
                code.append(lines[i])
                i += 1
            blocks.append(Block("code", "\n".join(code), lang=lang))
            i += 1
            continue

        if stripped == "":
            flush_paragraph()
            flush_list()
            i += 1
            continue

        if stripped == "---":
            flush_paragraph()
            flush_list()
            blocks.append(Block("hr", ""))
            i += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            flush_list()
            blocks.append(
                Block(
                    "heading",
                    heading.group(2).strip(),
                    level=len(heading.group(1)),
                )
            )
            i += 1
            continue

        bullet = re.match(r"^-\s+(.+)$", stripped)
        ordered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if bullet or ordered:
            flush_paragraph()
            kind = "ul" if bullet else "ol"
            if list_kind and list_kind != kind:
                flush_list()
            list_kind = kind
            list_items.append((bullet or ordered).group(1).strip())
            i += 1
            continue

        if stripped.startswith("> "):
            flush_paragraph()
            flush_list()
            blocks.append(Block("quote", stripped[2:].strip()))
            i += 1
            continue

        flush_list()
        paragraph.append(stripped)
        i += 1

    flush_paragraph()
    flush_list()
    return blocks


def inline_html(text: str) -> str:
    value = html.escape(text)
    value = re.sub(r"`([^`]+)`", r"<code>\1</code>", value)
    value = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", value)
    value = re.sub(
        r"(https?://[^\s<]+)",
        r'<a href="\1">\1</a>',
        value,
    )
    return value


def plain_inline(text: str) -> str:
    value = re.sub(r"`([^`]+)`", r"\1", text)
    value = re.sub(r"\*\*([^*]+)\*\*", r"\1", value)
    return value


def blocks_to_xhtml(blocks: list[Block], section_title: str) -> str:
    parts = [
        '<?xml version="1.0" encoding="utf-8"?>',
        '<!DOCTYPE html>',
        '<html xmlns="http://www.w3.org/1999/xhtml" lang="pt-BR">',
        "<head>",
        '<meta charset="utf-8"/>',
        f"<title>{html.escape(section_title)}</title>",
        '<link rel="stylesheet" type="text/css" href="style.css"/>',
        "</head><body>",
    ]
    for block in blocks:
        if block.kind == "heading":
            hid = slugify(block.text)
            level = min(block.level, 6)
            parts.append(
                f'<h{level} id="{hid}">{inline_html(block.text)}</h{level}>'
            )
        elif block.kind == "paragraph":
            parts.append(f"<p>{inline_html(block.text)}</p>")
        elif block.kind == "quote":
            parts.append(f"<blockquote>{inline_html(block.text)}</blockquote>")
        elif block.kind == "code":
            cls = f' class="language-{html.escape(block.lang)}"' if block.lang else ""
            parts.append(f"<pre><code{cls}>{html.escape(block.text)}</code></pre>")
        elif block.kind == "ul":
            parts.append("<ul>")
            for item in block.items or []:
                parts.append(f"<li>{inline_html(item)}</li>")
            parts.append("</ul>")
        elif block.kind == "ol":
            parts.append("<ol>")
            for item in block.items or []:
                parts.append(f"<li>{inline_html(item)}</li>")
            parts.append("</ol>")
        elif block.kind == "hr":
            parts.append("<hr/>")
    parts.append("</body></html>")
    return "\n".join(parts)


def split_sections(blocks: list[Block]) -> list[tuple[str, list[Block]]]:
    sections: list[tuple[str, list[Block]]] = []
    current_title = TITLE
    current: list[Block] = []

    for block in blocks:
        if block.kind == "heading" and block.level == 1 and current:
            sections.append((current_title, current))
            current_title = block.text
            current = [block]
        else:
            if block.kind == "heading" and block.level == 1:
                current_title = block.text
            current.append(block)

    if current:
        sections.append((current_title, current))
    return sections


def make_epub(blocks: list[Block]) -> None:
    tmp = WORK / "epub"
    if tmp.exists():
        shutil.rmtree(tmp)
    (tmp / "META-INF").mkdir(parents=True)
    oebps = tmp / "OEBPS"
    oebps.mkdir()

    (tmp / "mimetype").write_text("application/epub+zip", encoding="utf-8")
    (tmp / "META-INF" / "container.xml").write_text(
        """<?xml version="1.0" encoding="UTF-8"?>
<container version="1.0"
 xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
 <rootfiles>
  <rootfile full-path="OEBPS/content.opf"
   media-type="application/oebps-package+xml"/>
 </rootfiles>
</container>
""",
        encoding="utf-8",
    )
    (oebps / "style.css").write_text(
        """
body {
  font-family: serif;
  line-height: 1.45;
  margin: 0 5%;
}
h1, h2, h3 {
  font-family: sans-serif;
  line-height: 1.2;
}
h1 {
  margin-top: 1.2em;
  page-break-before: always;
}
pre {
  font-size: 0.82em;
  line-height: 1.25;
  white-space: pre-wrap;
  background: #f2f4f7;
  padding: 0.7em;
  border-left: 0.25em solid #2E74B5;
}
code {
  font-family: monospace;
}
blockquote {
  border-left: 0.25em solid #F0B84A;
  margin-left: 0;
  padding-left: 1em;
}
img.cover {
  width: 100%;
  height: auto;
}
""",
        encoding="utf-8",
    )
    shutil.copy(COVER, oebps / "cover.jpg")

    sections = split_sections(blocks)
    manifest_items = [
        '<item id="style" href="style.css" media-type="text/css"/>',
        '<item id="cover-image" href="cover.jpg" media-type="image/jpeg" properties="cover-image"/>',
        '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>',
        '<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>',
    ]
    spine_items = []
    nav_items = []
    ncx_points = []

    cover_xhtml = """<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml"
 xmlns:epub="http://www.idpf.org/2007/ops" lang="pt-BR">
<head><meta charset="utf-8"/><title>Capa</title>
<link rel="stylesheet" type="text/css" href="style.css"/></head>
<body><section epub:type="cover"><img class="cover" src="cover.jpg" alt="Capa"/></section></body>
</html>
"""
    (oebps / "cover.xhtml").write_text(cover_xhtml, encoding="utf-8")
    manifest_items.append(
        '<item id="cover-page" href="cover.xhtml" media-type="application/xhtml+xml"/>'
    )
    # KDP receives the cover JPG separately; keep cover metadata but do not
    # place a duplicate cover page in the reading order.

    for idx, (title, section_blocks) in enumerate(sections, start=1):
        filename = f"section-{idx:02d}.xhtml"
        item_id = f"section-{idx:02d}"
        (oebps / filename).write_text(
            blocks_to_xhtml(section_blocks, title),
            encoding="utf-8",
        )
        manifest_items.append(
            f'<item id="{item_id}" href="{filename}" media-type="application/xhtml+xml"/>'
        )
        spine_items.append(f'<itemref idref="{item_id}"/>')
        nav_items.append(f'<li><a href="{filename}">{html.escape(title)}</a></li>')
        ncx_points.append(
            f"""<navPoint id="navPoint-{idx}" playOrder="{idx}">
  <navLabel><text>{xml_escape(title)}</text></navLabel>
  <content src="{filename}"/>
</navPoint>"""
        )

    (oebps / "nav.xhtml").write_text(
        f"""<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml"
 xmlns:epub="http://www.idpf.org/2007/ops" lang="pt-BR">
<head><meta charset="utf-8"/><title>Sumário</title>
<link rel="stylesheet" type="text/css" href="style.css"/></head>
<body>
<nav epub:type="toc" id="toc"><h1>Sumário</h1><ol>
{chr(10).join(nav_items)}
</ol></nav>
</body></html>
""",
        encoding="utf-8",
    )

    (oebps / "toc.ncx").write_text(
        f"""<?xml version="1.0" encoding="UTF-8"?>
<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">
<head>
 <meta name="dtb:uid" content="robo-trade-go-binance"/>
 <meta name="dtb:depth" content="1"/>
 <meta name="dtb:totalPageCount" content="0"/>
 <meta name="dtb:maxPageNumber" content="0"/>
</head>
<docTitle><text>{xml_escape(TITLE)}</text></docTitle>
<navMap>
{chr(10).join(ncx_points)}
</navMap>
</ncx>
""",
        encoding="utf-8",
    )

    modified = (
        datetime.now(timezone.utc)
        .replace(microsecond=0)
        .isoformat()
        .replace("+00:00", "Z")
    )
    (oebps / "content.opf").write_text(
        f"""<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf"
 version="3.0" unique-identifier="bookid">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
 <dc:identifier id="bookid">robo-trade-go-binance</dc:identifier>
 <dc:title>{xml_escape(TITLE)}</dc:title>
 <dc:creator>{xml_escape(AUTHOR)}</dc:creator>
 <dc:language>{LANG}</dc:language>
 <dc:description>{xml_escape(SUBTITLE)}</dc:description>
 <dc:rights>{xml_escape(COPYRIGHT)}</dc:rights>
 <dc:publisher>{xml_escape(AUTHOR)}</dc:publisher>
 <meta property="dcterms:modified">{modified}</meta>
 <meta name="cover" content="cover-image"/>
</metadata>
<manifest>
 {chr(10).join(manifest_items)}
</manifest>
<spine toc="ncx">
 {chr(10).join(spine_items)}
</spine>
</package>
""",
        encoding="utf-8",
    )

    if EPUB.exists():
        EPUB.unlink()
    with zipfile.ZipFile(EPUB, "w") as zf:
        zf.write(tmp / "mimetype", "mimetype", compress_type=zipfile.ZIP_STORED)
        for path in sorted(tmp.rglob("*")):
            if path.name == "mimetype" or path.is_dir():
                continue
            zf.write(path, path.relative_to(tmp), compress_type=zipfile.ZIP_DEFLATED)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def make_docx(blocks: list[Block]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.25

    code_style = styles.add_style("CodeBlock", 1)
    code_style.font.name = "Courier New"
    code_style.font.size = Pt(8.5)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(8)
    code_style.paragraph_format.line_spacing = 1.05

    first_h1 = True
    for block in blocks:
        if block.kind == "heading":
            if block.level == 1 and not first_h1:
                doc.add_page_break()
            first_h1 = False if block.level == 1 else first_h1
            level = min(block.level, 3)
            doc.add_heading(plain_inline(block.text), level=level)
        elif block.kind == "paragraph":
            doc.add_paragraph(plain_inline(block.text))
        elif block.kind == "quote":
            p = doc.add_paragraph(plain_inline(block.text))
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(4)
            shade_paragraph(p, "F8F3E8")
        elif block.kind == "code":
            p = doc.add_paragraph(style="CodeBlock")
            p.add_run(block.text)
            shade_paragraph(p, "F2F4F7")
        elif block.kind == "ul":
            for item in block.items or []:
                doc.add_paragraph(plain_inline(item), style="List Bullet")
        elif block.kind == "ol":
            for item in block.items or []:
                doc.add_paragraph(plain_inline(item), style="List Number")
        elif block.kind == "hr":
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.LINE)

    doc.core_properties.title = TITLE
    doc.core_properties.author = AUTHOR
    doc.core_properties.language = LANG
    doc.save(DOCX)


def make_pdf(blocks: list[Block]) -> None:
    page_size = (5 * inch, 8 * inch)
    doc = SimpleDocTemplate(
        str(PDF),
        pagesize=page_size,
        rightMargin=0.55 * rl_inch,
        leftMargin=0.55 * rl_inch,
        topMargin=0.58 * rl_inch,
        bottomMargin=0.58 * rl_inch,
        title=TITLE,
        author=AUTHOR,
    )

    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "BookBody",
        parent=styles["BodyText"],
        fontName="Times-Roman",
        fontSize=11.1,
        leading=16.1,
        spaceAfter=7,
        alignment=TA_LEFT,
    )
    h1 = ParagraphStyle(
        "BookH1",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0B3B4A"),
        spaceAfter=14,
        alignment=TA_LEFT,
    )
    h2 = ParagraphStyle(
        "BookH2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=17,
        textColor=colors.HexColor("#1F4D78"),
        spaceBefore=8,
        spaceAfter=7,
    )
    h3 = ParagraphStyle(
        "BookH3",
        parent=styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#1F4D78"),
        spaceBefore=6,
        spaceAfter=5,
    )
    code = ParagraphStyle(
        "Code",
        parent=styles["Code"],
        fontName="Courier",
        fontSize=7.4,
        leading=10.0,
        backColor=colors.HexColor("#F2F4F7"),
        borderColor=colors.HexColor("#D8DEE6"),
        borderWidth=0.4,
        borderPadding=5,
        spaceAfter=8,
    )
    quote = ParagraphStyle(
        "Quote",
        parent=body,
        leftIndent=14,
        textColor=colors.HexColor("#324B55"),
        backColor=colors.HexColor("#F8F3E8"),
        borderPadding=5,
    )
    center = ParagraphStyle(
        "Center",
        parent=body,
        alignment=TA_CENTER,
    )

    story = [
        PdfImage(str(COVER), width=3.9 * rl_inch, height=6.24 * rl_inch),
        PageBreak(),
    ]

    first_h1 = True
    for block in blocks:
        if block.kind == "heading":
            if block.level == 1 and not first_h1:
                story.append(PageBreak())
            if block.level == 1:
                first_h1 = False
                story.append(Paragraph(inline_html(block.text), h1))
            elif block.level == 2:
                story.append(Paragraph(inline_html(block.text), h2))
            else:
                story.append(Paragraph(inline_html(block.text), h3))
        elif block.kind == "paragraph":
            story.append(Paragraph(inline_html(block.text), body))
        elif block.kind == "quote":
            story.append(Paragraph(inline_html(block.text), quote))
        elif block.kind == "code":
            wrapped = []
            for line in block.text.splitlines():
                wrapped.extend(textwrap.wrap(line, width=54) or [""])
            story.append(Preformatted("\n".join(wrapped), code))
        elif block.kind in {"ul", "ol"}:
            items = [
                ListItem(Paragraph(inline_html(item), body), leftIndent=12)
                for item in (block.items or [])
            ]
            story.append(
                ListFlowable(
                    items,
                    bulletType="bullet" if block.kind == "ul" else "1",
                    start="circle" if block.kind == "ul" else "1",
                    leftIndent=18,
                )
            )
            story.append(Spacer(1, 3))
        elif block.kind == "hr":
            story.append(Spacer(1, 10))
            story.append(Paragraph("* * *", center))
            story.append(Spacer(1, 10))

    def footer(canvas, doc_obj):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#69777D"))
        canvas.drawCentredString(
            page_size[0] / 2,
            0.33 * rl_inch,
            str(canvas.getPageNumber()),
        )
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def make_print_pdf(blocks: list[Block]) -> int:
    """Miolo de impressão 6x9 P&B, margens espelhadas, sem capa. Retorna páginas."""
    width, height = TRIM_6X9[0] * inch, TRIM_6X9[1] * inch
    inside, outside, top_bottom = 0.625 * inch, 0.5 * inch, 0.6 * inch
    frame_w = width - inside - outside
    frame_h = height - 2 * top_bottom
    recto = Frame(inside, top_bottom, frame_w, frame_h, id="recto")  # gutter à esq.
    verso = Frame(outside, top_bottom, frame_w, frame_h, id="verso")  # gutter à dir.

    def footer(canvas, doc_obj):
        start = getattr(doc_obj, "_body_start", None)
        if start and canvas.getPageNumber() >= start:
            canvas.saveState()
            canvas.setFont("Helvetica", 8)
            canvas.setFillColor(colors.HexColor("#444444"))
            canvas.drawCentredString(width / 2, 0.32 * inch, str(canvas.getPageNumber()))
            canvas.restoreState()

    class MirrorDoc(BaseDocTemplate):
        def handle_pageBegin(self):
            self._handle_pageBegin()
            self.handle_nextPageTemplate("verso" if self.page % 2 else "recto")

    doc = MirrorDoc(
        str(PRINT_PDF), pagesize=(width, height), title=TITLE, author=AUTHOR
    )
    doc.addPageTemplates(
        [
            PageTemplate(id="recto", frames=[recto], onPage=footer),
            PageTemplate(id="verso", frames=[verso], onPage=footer),
        ]
    )

    class BodyStart(Flowable):
        def wrap(self, *_):
            return (0, 0)

        def draw(self):
            if not getattr(doc, "_body_start", None):
                doc._body_start = doc.page

    body = ParagraphStyle(
        "PrintBody", fontName="Times-Roman", fontSize=10.5, leading=14.5,
        spaceAfter=6, alignment=TA_JUSTIFY,
    )
    h1 = ParagraphStyle(
        "PrintH1", fontName="Helvetica-Bold", fontSize=17, leading=21,
        textColor=colors.black, spaceBefore=18, spaceAfter=14,
    )
    h2 = ParagraphStyle(
        "PrintH2", fontName="Helvetica-Bold", fontSize=12.5, leading=16,
        textColor=colors.HexColor("#1A1A1A"), spaceBefore=9, spaceAfter=6,
    )
    h3 = ParagraphStyle(
        "PrintH3", fontName="Helvetica-Bold", fontSize=11, leading=14,
        textColor=colors.HexColor("#333333"), spaceBefore=6, spaceAfter=4,
    )
    code = ParagraphStyle(
        "PrintCode", fontName="Courier", fontSize=8, leading=10.8,
        backColor=colors.HexColor("#F0F0F0"), borderColor=colors.HexColor("#CCCCCC"),
        borderWidth=0.4, borderPadding=5, spaceAfter=8,
    )
    quote = ParagraphStyle(
        "PrintQuote", parent=body, leftIndent=14, fontName="Times-Italic",
        textColor=colors.HexColor("#333333"),
    )
    center = ParagraphStyle("PrintCenter", parent=body, alignment=TA_CENTER)

    story: list = []
    h1_count = 0
    body_started = False
    for block in blocks:
        if block.kind == "heading":
            if block.level == 1:
                h1_count += 1
                if h1_count >= 2:  # segundo h1 = início dos capítulos (corpo)
                    story.append(PageBreak())
                    if not body_started:
                        story.append(BodyStart())
                        body_started = True
                story.append(Paragraph(inline_html(block.text), h1))
            elif block.level == 2:
                story.append(Paragraph(inline_html(block.text), h2))
            else:
                story.append(Paragraph(inline_html(block.text), h3))
        elif block.kind == "paragraph":
            story.append(Paragraph(inline_html(block.text), body))
        elif block.kind == "quote":
            story.append(Paragraph(inline_html(block.text), quote))
        elif block.kind == "code":
            wrapped = []
            for line in block.text.splitlines():
                wrapped.extend(textwrap.wrap(line, width=70) or [""])
            story.append(Preformatted("\n".join(wrapped), code))
        elif block.kind in {"ul", "ol"}:
            items = [
                ListItem(Paragraph(inline_html(item), body), leftIndent=12)
                for item in (block.items or [])
            ]
            story.append(
                ListFlowable(
                    items,
                    bulletType="bullet" if block.kind == "ul" else "1",
                    start="circle" if block.kind == "ul" else "1",
                    leftIndent=18,
                )
            )
            story.append(Spacer(1, 3))
        elif block.kind == "hr":
            if body_started:
                story.append(Spacer(1, 10))
                story.append(Paragraph("* * *", center))
                story.append(Spacer(1, 10))
            else:
                story.append(PageBreak())  # separa páginas do front matter

    doc.build(story)
    return doc.page


def make_code_zip() -> None:
    if CODE_ZIP.exists():
        CODE_ZIP.unlink()
    base = shutil.make_archive(str(CODE_ZIP.with_suffix("")), "zip", CODE_DIR)
    Path(base).replace(CODE_ZIP)


def main() -> None:
    markdown = MD.read_text(encoding="utf-8")
    blocks = parse_markdown(markdown)
    make_cover()
    make_epub(blocks)
    make_docx(blocks)
    make_pdf(blocks)
    print_pages = make_print_pdf(blocks)  # miolo antes da capa (define a lombada)
    make_print_cover(print_pages)
    make_code_zip()
    print(f"cover:       {COVER}")
    print(f"epub:        {EPUB}")
    print(f"docx:        {DOCX}")
    print(f"pdf (revisão): {PDF}")
    print(f"print miolo:   {PRINT_PDF} ({print_pages} páginas)")
    print(f"print capa:    {PRINT_COVER}")
    print(f"zip:         {CODE_ZIP}")


if __name__ == "__main__":
    main()
